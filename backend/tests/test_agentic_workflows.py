import os
import tempfile
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from docx import Document as DocxDocument
from openpyxl import load_workbook
from app.main import app
from app.core.database import init_db
from app.seed.demo_data import seed_demo_data
from app.services.agent.python_sandbox import execute_sandboxed_math
from app.services.agent.deliverable_generator import generate_approval_docx, generate_audit_xlsx

@pytest_asyncio.fixture(autouse=True)
async def setup_db():
    await init_db()
    await seed_demo_data()

@pytest.mark.asyncio
async def test_python_sandbox_math_execution():
    inputs = {"current_pressure_bar": 48.5, "threshold_pressure_bar": 45.0}
    res = execute_sandboxed_math("", inputs)
    assert res["success"] is True
    assert res["exceedance_bar"] == 3.5
    assert res["exceedance_pct"] == 7.78
    assert res["is_exceeded"] is True

@pytest.mark.asyncio
async def test_physical_docx_generation():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        file_path = generate_approval_docx(file_path=tmp_path, exceedance_pct=7.78)
        assert os.path.exists(file_path)
        assert os.path.getsize(file_path) > 1000

        doc = DocxDocument(file_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])
        assert "CRITICAL GOVERNANCE ALERT" in full_text
        assert "48.5 BAR" in full_text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

@pytest.mark.asyncio
async def test_physical_xlsx_generation():
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        file_path = generate_audit_xlsx(file_path=tmp_path)
        assert os.path.exists(file_path)
        assert os.path.getsize(file_path) > 1000

        wb = load_workbook(file_path)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Timestamp"
        assert ws.cell(row=2, column=2).value == "P-102A"
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

@pytest.mark.asyncio
async def test_workflow_execution_api():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # Fetch workflow templates
        wf_res = await ac.get("/api/v1/workflows/templates")
        assert wf_res.status_code == 200
        wfs = wf_res.json()
        assert len(wfs) >= 1
        wf_id = wfs[0]["id"]

        # Execute workflow
        exec_payload = {"prompt": "Run Pump P-102A inspection review workflow."}
        exec_res = await ac.post(f"/api/v1/workflows/{wf_id}/execute", json=exec_payload)
        assert exec_res.status_code == 200
        run_data = exec_res.json()

        assert run_data["status"] == "COMPLETED"
        assert run_data["total_steps"] == 5
        assert "Step 1/5" in run_data["output_summary"]
        assert "Step 5/5" in run_data["output_summary"]

@pytest.mark.asyncio
async def test_deliverable_approval_and_download_api():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # List deliverables
        res = await ac.get("/api/v1/deliverables")
        assert res.status_code == 200
        delivs = res.json()
        assert len(delivs) >= 1
        deliv = delivs[0]
        deliv_id = deliv["id"]

        # Approve deliverable
        appr_res = await ac.post(f"/api/v1/deliverables/{deliv_id}/approve")
        assert appr_res.status_code == 200
        assert appr_res.json()["approval_status"] == "APPROVED"

        # Download deliverable file
        down_res = await ac.get(f"/api/v1/deliverables/{deliv_id}/download")
        assert down_res.status_code == 200
        assert len(down_res.content) > 1000

        # Verify audit logs
        audit_res = await ac.get("/api/v1/audit/events")
        assert audit_res.status_code == 200
        actions = [a["action"] for a in audit_res.json()]
        assert "DELIVERABLE_APPROVAL" in actions
