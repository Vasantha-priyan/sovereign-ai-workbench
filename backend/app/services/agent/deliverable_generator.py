import os
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.core.config import settings

def generate_approval_docx(
    file_path: str,
    title: str = "Emergency Shutdown Approval Note",
    equipment_id: str = "Pump P-102A",
    pressure_bar: float = 48.5,
    threshold_bar: float = 45.0,
    exceedance_pct: float = 7.78
) -> str:
    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    doc = Document()

    # Title Banner
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.color.rgb = RGBColor(180, 0, 0) # Red
        run.font.size = Pt(20)

    # Subtitle
    sub = doc.add_paragraph(f"Refinery Unit 4B — Confidential Engineering Governance Note\nDate: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.italic = True
    sub.runs[0].font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Executive Summary Callout
    p1 = doc.add_paragraph()
    r1 = p1.add_run("CRITICAL GOVERNANCE ALERT: DISCHARGE PRESSURE EXCEEDANCE\n")
    r1.bold = True
    r1.font.color.rgb = RGBColor(200, 0, 0)
    p1.add_run(
        f"During routine inspection, target equipment '{equipment_id}' registered an operational discharge pressure of "
        f"{pressure_bar} BAR, exceeding the mandatory Safety SOP threshold limit of {threshold_bar} BAR by +{exceedance_pct}%."
    )

    # Table of Technical Metrics
    doc.add_heading("1. Operational Parameters & Verification", level=2)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"

    metrics = [
        ("Equipment Target", equipment_id),
        ("Measured Pressure", f"{pressure_bar} BAR"),
        ("Safety SOP Limit", f"{threshold_bar} BAR"),
        ("Exceedance Delta", f"+{round(pressure_bar - threshold_bar, 2)} BAR (+{exceedance_pct}%)"),
        ("Required Action", "Immediate Load Transfer & Isolation")
    ]

    for i, (k, v) in enumerate(metrics):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        row.cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Corrective Action Plan
    doc.add_heading("2. Mandatory Corrective Action Plan", level=2)
    doc.add_paragraph("1. Initiate immediate electrical and hydraulic isolation of Pump P-102A.", style="List Bullet")
    doc.add_paragraph("2. Transfer operational load to secondary backup pump P-102B.", style="List Bullet")
    doc.add_paragraph("3. Conduct full non-destructive testing (NDT) inspection of discharge valve assembly.", style="List Bullet")

    # Signature Block
    doc.add_heading("3. Human-in-the-Loop Sign-Off", level=2)
    sig_p = doc.add_paragraph(
        "Approved by: _________________________\n"
        "Rajesh Verma — Senior Operations Engineer\n"
        "Status: VERIFIED & APPROVED FOR EXECUTION"
    )
    sig_p.runs[0].bold = True

    doc.save(file_path)
    return file_path

def generate_audit_xlsx(
    file_path: str,
    title: str = "Equipment Inspection Audit Summary"
) -> str:
    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Audit Log"

    # Header Row Styling
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ["Timestamp", "Equipment ID", "Unit", "Measured Pressure (BAR)", "SOP Limit (BAR)", "Status", "Action Taken"]
    ws.append(headers)

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data Rows
    rows_data = [
        ("2026-08-27T08:00:00Z", "P-102A", "Refinery 4B", 48.5, 45.0, "ALERT", "Emergency Isolation Triggered"),
        ("2026-08-27T09:00:00Z", "P-102A", "Refinery 4B", 48.8, 45.0, "ALERT", "Load Transferred to P-102B"),
        ("2026-08-27T10:00:00Z", "P-102B", "Refinery 4B", 41.2, 45.0, "NORMAL", "Backup Operating Normally")
    ]

    for r in rows_data:
        ws.append(list(r))

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(file_path)
    return file_path
