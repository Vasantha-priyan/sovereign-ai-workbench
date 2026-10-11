import os
from datetime import datetime
from app.core.config import settings

class DocGenService:
    def create_docx_approval_note(
        self,
        title: str,
        subject: str,
        findings: list[str],
        sop_references: list[str],
        recommendation: str,
        output_filename: str = "Approval_Note.docx"
    ) -> str:
        output_path = os.path.join(settings.DELIVERABLES_DIR, output_filename)
        
        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = docx.Document()

            title_p = doc.add_paragraph()
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title_p.add_run("OFFICIAL ENTERPRISE APPROVAL NOTE")
            run.font.name = "Arial"
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 34, 64)

            sub_p = doc.add_paragraph()
            sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_sub = sub_p.add_run("CONFIDENTIAL — FOR INTERNAL USE ONLY")
            run_sub.font.name = "Arial"
            run_sub.font.size = Pt(10)
            run_sub.font.bold = True
            run_sub.font.color.rgb = RGBColor(180, 40, 40)

            doc.add_paragraph("=" * 60)

            meta_table = doc.add_table(rows=4, cols=2)
            labels = [
                ("Date:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                ("Subject:", subject),
                ("Originator:", "AI Workbench ReAct Agent / Plant Operations"),
                ("Status:", "PENDING HUMAN REVIEW & APPROVAL")
            ]
            for idx, (label, val) in enumerate(labels):
                row_cells = meta_table.rows[idx].cells
                row_cells[0].paragraphs[0].add_run(label).bold = True
                row_cells[1].paragraphs[0].add_run(val)

            doc.add_paragraph()
            doc.add_heading("1. Executive Summary & Findings", level=1)
            for f in findings:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f)

            doc.add_heading("2. Grounded SOP & Safety Policy Compliance", level=1)
            for ref in sop_references:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(ref)

            doc.add_heading("3. Official AI Recommendation", level=1)
            p_rec = doc.add_paragraph()
            p_rec.add_run(recommendation).italic = True

            doc.add_paragraph()
            sign_p = doc.add_paragraph()
            sign_p.add_run("AI-Generated Draft Recommendation — Human Approval Required.\n\n").bold = True
            sign_p.add_run("Approver Signature: _______________________      Date: ______________")

            doc.save(output_path)
            return output_path
        except Exception:
            # Fallback document creation
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"OFFICIAL APPROVAL NOTE\n{title}\n\nSubject: {subject}\nDate: {datetime.now()}\n\nFINDINGS:\n" + "\n".join(findings) + f"\n\nRECOMMENDATION:\n{recommendation}")
            return output_path

    def create_xlsx_summary(
        self,
        title: str,
        headers: list[str],
        data_rows: list[list],
        output_filename: str = "Inspection_Summary.xlsx"
    ) -> str:
        output_path = os.path.join(settings.DELIVERABLES_DIR, output_filename)
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Summary Audit"

            ws.merge_cells("A1:E1")
            title_cell = ws["A1"]
            title_cell.value = title.upper()
            title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.append([])
            ws.append(headers)

            for row in data_rows:
                ws.append(row)

            wb.save(output_path)
            return output_path
        except Exception:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"{title}\n" + " | ".join(headers) + "\n")
                for r in data_rows:
                    f.write(" | ".join(str(x) for x in r) + "\n")
            return output_path

    def create_pptx_slides(
        self,
        title: str,
        subtitle: str,
        slides_content: list[dict],
        output_filename: str = "Executive_Briefing.pptx"
    ) -> str:
        output_path = os.path.join(settings.DELIVERABLES_DIR, output_filename)
        try:
            from pptx import Presentation
            prs = Presentation()

            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            slide.shapes.title.text = title
            slide.placeholders[1].text = subtitle

            bullet_slide_layout = prs.slide_layouts[1]
            for item in slides_content:
                s = prs.slides.add_slide(bullet_slide_layout)
                s.shapes.title.text = item.get("slide_title", "Summary")
                tf = s.placeholders[1].text_frame
                for point in item.get("bullet_points", []):
                    p = tf.add_paragraph()
                    p.text = point

            prs.save(output_path)
            return output_path
        except Exception:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"PPTX BRIEFING:\n{title}\n{subtitle}")
            return output_path

docgen_service = DocGenService()
